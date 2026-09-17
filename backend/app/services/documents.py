import re
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from pypdf import PdfReader
from docx import Document
from ..schemas import PersonaData

HEADINGS = {
    "education": "education",
    "教育经历": "education",
    "教育背景": "education",
    "experience": "experience",
    "work experience": "experience",
    "工作经历": "experience",
    "professional experience": "experience",
    "skills": "skills",
    "技能": "skills",
    "sectors": "sectors",
    "金融领域": "sectors",
    "career goals": "career_goals",
    "职业目标": "career_goals",
    "target regions": "target_regions",
    "目标地区": "target_regions",
    "target roles": "target_roles",
    "目标机构或职位": "target_roles",
    "contact purpose": "contact_purpose",
    "联系目的": "contact_purpose",
}


def extract_text(content: bytes, extension: str):
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("This file is not a valid PDF. Use manual entry instead.")
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ValueError(
                "Password-protected PDF is unsupported. Upload an unlocked copy or enter your background manually."
            )
        if len(reader.pages) > 30:
            raise ValueError("Resume exceeds the 30-page limit. Upload a shorter file.")
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
    elif extension == ".docx":
        try:
            with ZipFile(BytesIO(content)) as z:
                if (
                    sum(i.file_size for i in z.infolist()) > 25 * 1024 * 1024
                    or len(z.infolist()) > 1000
                ):
                    raise ValueError("DOCX expands beyond the safe size limit.")
                if "word/document.xml" not in z.namelist():
                    raise ValueError("Invalid DOCX document.")
            doc = Document(BytesIO(content))
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [c.text for t in doc.tables for row in t.rows for c in row.cells]
            )
        except BadZipFile:
            raise ValueError("This file is not a valid DOCX. Use manual entry instead.")
    else:
        raise ValueError("Only text-based PDF and DOCX resumes are supported.")
    if len(re.sub(r"\s", "", text)) < 30:
        raise ValueError(
            "No usable text was found. Scanned/image-only resumes are unsupported (no OCR). Upload a text-based PDF/DOCX or use manual entry."
        )
    if len(text) > 50000:
        raise ValueError("Extracted text is too long (maximum 50,000 characters).")
    return text


def extract_sections(text):
    data = PersonaData().model_dump()
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    first = lines[0] if lines else ""
    if 1 < len(first) < 80 and not any(
        x in first.lower() for x in ("@", "resume", "curriculum", "简历", "http")
    ):
        data["name"] = first
    field = None
    for line in lines[1:]:
        parts = re.split(r"[:：]", line, maxsplit=1)
        key = HEADINGS.get(parts[0].lower().strip())
        if key:
            field = key
            if len(parts) > 1:
                data[field] += parts[1].strip() + "\n"
        elif field:
            data[field] += line + "\n"
    return {k: v.strip() for k, v in data.items()}


# Upload storage is durable before the request returns. A single process consumes
# queued records; provider work never holds a database transaction or persona lock.
from hashlib import sha256
from pathlib import Path
from threading import Event, Lock, Thread
from uuid import uuid4
from fastapi import HTTPException
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError
from ..config import settings
from ..db import Session
from ..models import UploadedDocument, DocumentFile, now
from .file_storage import read_content
from ..providers.ai import CompatibleAI, MockAI, resolve_model

PARSE_PROMPT_VERSION = "resume-parse-v1"


def document_response(document):
    return {
        "id": document.id,
        "document_id": document.id,
        "original_name": document.original_name,
        "status": document.status,
        "data": document.parsed_data,
        "error": document.error,
        "extracted_text": document.extracted_text,
        "persona_id": document.persona_id,
        "mode": document.parse_mode,
        "provider": document.parse_provider,
        "model": document.parse_model,
        "prompt_version": document.parse_prompt_version,
        "created_at": document.created_at,
        "started_at": document.parse_started_at,
        "completed_at": document.parse_completed_at,
        "notice": "Review every field before saving. Parsing does not change an existing sender profile.",
    }


def cached_document(repo, digest, mode, provider, model, exclude_id=None):
    conditions = [
        UploadedDocument.content_hash == digest,
        UploadedDocument.parse_mode == mode,
        UploadedDocument.parse_provider == provider,
        UploadedDocument.parse_model == model,
        UploadedDocument.parse_prompt_version == PARSE_PROMPT_VERSION,
        UploadedDocument.status.in_(["queued", "processing", "parsed"]),
    ]
    if exclude_id:
        conditions.append(UploadedDocument.id != exclude_id)
    found = repo.all(UploadedDocument, *conditions)
    return found[0] if found else None


def queue_document(repo, file):
    name = Path(file.filename or "resume").name[:255]
    ext = Path(name).suffix.lower()
    if ext not in (".pdf", ".docx"):
        raise HTTPException(
            415,
            "Only text-based PDF or DOCX files are supported. You can enter your background manually.",
        )
    content = file.file.read(8 * 1024 * 1024 + 1)
    if len(content) > 8 * 1024 * 1024:
        raise HTTPException(
            413, "File exceeds 8 MB. Use a smaller resume or manual entry."
        )
    digest, mode, provider, model = (
        sha256(content).hexdigest(),
        settings.ai_mode,
        settings.ai_provider,
        resolve_model(),
    )
    cached = cached_document(repo, digest, mode, provider, model)
    if cached:
        return cached
    if (
        len(
            repo.all(
                UploadedDocument, UploadedDocument.status.in_(["queued", "processing"])
            )
        )
        >= 5
    ):
        raise HTTPException(
            429,
            "Five uploads are still being parsed. Wait for a result before uploading another file.",
        )
    directory = Path(settings.upload_dir).resolve() / repo.workspace_id
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    key = str(uuid4()) + ext
    path = directory / key
    path.write_bytes(content)
    path.chmod(0o600)
    try:
        with repo.session.begin_nested():
            document = repo.add(
                UploadedDocument,
                original_name=name,
                storage_key=key,
                status="queued",
                extracted_text="",
                content_hash=digest,
                parse_mode=mode,
                parse_provider=provider,
                parse_model=model,
                parse_prompt_version=PARSE_PROMPT_VERSION,
            )
            repo.session.add(DocumentFile(document_id=document.id, byte_size=len(content), content=content))
            repo.session.flush()
            return document
    except IntegrityError:
        path.unlink(missing_ok=True)
        cached = cached_document(repo, digest, mode, provider, model)
        if cached:
            return cached
        raise


def retry_document(repo, document):
    if document.status != "failed":
        return document
    mode, provider, model = settings.ai_mode, settings.ai_provider, resolve_model()
    if document.content_hash:
        cached = cached_document(
            repo, document.content_hash, mode, provider, model, document.id
        )
        if cached:
            return cached
    if (
        len(
            repo.all(
                UploadedDocument, UploadedDocument.status.in_(["queued", "processing"])
            )
        )
        >= 5
    ):
        raise HTTPException(
            429,
            "Five uploads are still being parsed. Wait before retrying another file.",
        )
    with repo.session.no_autoflush:
        repo.session.execute(
            update(UploadedDocument)
            .where(
                UploadedDocument.id == document.id,
                UploadedDocument.workspace_id == repo.workspace_id,
                UploadedDocument.status == "failed",
            )
            .values(
                status="queued",
                error=None,
                parsed_data=None,
                parse_mode=mode,
                parse_provider=provider,
                parse_model=model,
                parse_prompt_version=PARSE_PROMPT_VERSION,
                parse_started_at=None,
                parse_completed_at=None,
            )
            .execution_options(synchronize_session=False)
        )
    repo.session.expire(document)
    repo.session.refresh(document)
    return document


def process_document(document_id):
    with Session() as db:
        claimed = db.execute(
            update(UploadedDocument)
            .where(
                UploadedDocument.id == document_id, UploadedDocument.status == "queued"
            )
            .values(status="processing", parse_started_at=now(), error=None)
        )
        if not claimed.rowcount:
            db.rollback()
            return False
        document = db.get(UploadedDocument, document_id)
        path = (
            Path(settings.upload_dir).resolve()
            / document.workspace_id
            / document.storage_key
        )
        mode, provider_name, model = (
            document.parse_mode,
            document.parse_provider,
            document.parse_model,
        )
        prompt_version, digest = document.parse_prompt_version, document.content_hash
        # Keep the original available even when a free host discards its local disk.
        try:
            stored_content = read_content(db, document)
        except HTTPException:
            stored_content = None
        db.commit()
    extracted, data, error = "", None, None
    try:
        if (
            prompt_version != PARSE_PROMPT_VERSION
            or mode != settings.ai_mode
            or provider_name != settings.ai_provider
        ):
            raise HTTPException(
                409,
                "The parser configuration changed while this upload was queued. Retry parsing with the current configuration.",
            )
        if stored_content is None:
            raise ValueError("Original file is unavailable. Please upload it again.")
        content = stored_content
        if digest and sha256(content).hexdigest() != digest:
            raise ValueError(
                "The stored file changed after upload. Upload the original file again."
            )
        extracted = extract_text(content, path.suffix.lower())
        provider = MockAI() if mode == "mock" else CompatibleAI()
        result = provider.complete("parse", {"text": extracted, "_model": model})
        if not isinstance(result, dict) or not isinstance(result.get("data"), dict):
            raise ValueError(
                "The AI parser returned invalid data. Retry parsing or enter the background manually."
            )
        data = PersonaData.model_validate(result["data"]).model_dump()
        if not any(v.strip() for v in data.values()):
            raise ValueError(
                "The parser found no usable profile fields. Review the extracted text or use manual entry."
            )
        status = "parsed"
    except Exception as exc:
        status = "failed"
        if isinstance(exc, HTTPException):
            error = str(exc.detail)
        elif isinstance(exc, ValueError) and not hasattr(exc, "errors"):
            error = str(exc)
        else:
            error = "Unable to parse this document. It may be malformed or unsupported. Retry parsing or use manual entry."
        if settings.ai_api_key:
            error = error.replace(settings.ai_api_key, "[redacted]")
    with Session() as db:
        db.execute(
            update(UploadedDocument)
            .where(
                UploadedDocument.id == document_id,
                UploadedDocument.status == "processing",
            )
            .values(
                status=status,
                parsed_data=data,
                extracted_text=extracted,
                error=error,
                parse_completed_at=now(),
            )
        )
        db.commit()
    return True


def recover_document_jobs():
    with Session() as db:
        db.execute(
            update(UploadedDocument)
            .where(UploadedDocument.status == "processing")
            .values(
                status="failed",
                parse_completed_at=now(),
                error="The server restarted while parsing. Your file is saved. Retry parsing explicitly; the previous AI call may have completed.",
            )
        )
        db.commit()


_document_worker = None
_document_guard = Lock()
_document_wake = Event()
_document_stop = Event()


def wake_document_worker():
    _document_wake.set()


def _document_loop():
    while not _document_stop.is_set():
        try:
            with Session() as db:
                document_id = db.scalar(
                    select(UploadedDocument.id)
                    .where(UploadedDocument.status == "queued")
                    .order_by(UploadedDocument.created_at)
                    .limit(1)
                )
            if document_id:
                process_document(document_id)
                continue
        except Exception:
            pass
        _document_wake.wait(1)
        _document_wake.clear()


def start_document_worker():
    global _document_worker
    with _document_guard:
        if _document_worker and _document_worker.is_alive():
            return
        recover_document_jobs()
        _document_stop.clear()
        _document_wake.clear()
        _document_worker = Thread(
            target=_document_loop, name="document-parsing", daemon=True
        )
        _document_worker.start()


def stop_document_worker():
    _document_stop.set()
    _document_wake.set()
    if _document_worker:
        _document_worker.join(timeout=5)
