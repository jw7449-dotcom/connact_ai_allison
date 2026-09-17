from io import BytesIO
from threading import Event
from concurrent.futures import ThreadPoolExecutor
import time
import pytest
from docx import Document
from fastapi import HTTPException
from app.db import Session
from app.models import Workspace, UploadedDocument
from app.providers.ai import MockAI
from app.services.documents import (
    process_document,
    recover_document_jobs,
    start_document_worker,
    stop_document_worker,
)


@pytest.fixture
def paused_documents(client):
    stop_document_worker()
    yield client
    stop_document_worker()


def resume_bytes(name="Jordan Example"):
    doc = Document()
    doc.add_paragraph(name)
    doc.add_paragraph("Education")
    doc.add_paragraph("Example University, Bachelor of Finance")
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python, valuation, financial modeling")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def queue(c, content=None, filename="resume.docx"):
    r = c.post(
        "/api/documents/jobs", files={"file": (filename, content or resume_bytes())}
    )
    assert r.status_code == 202, r.text
    return r.json()


def test_uploaded_parse_is_persistent_and_reupload_reuses_paid_result(
    paused_documents, monkeypatch
):
    c = paused_documents
    original, calls = MockAI.complete, []

    def count(self, task, data):
        calls.append(task)
        return original(self, task, data)

    monkeypatch.setattr(MockAI, "complete", count)
    content = resume_bytes()
    job = queue(c, content)
    assert job["status"] == "queued" and job["data"] is None
    assert queue(c, content)["id"] == job["id"]
    assert process_document(job["id"])
    assert not process_document(job["id"])
    saved = c.get(f'/api/documents/{job["id"]}/status').json()
    assert saved["status"] == "parsed" and saved["data"]["name"] == "Jordan Example"
    assert "Example University" in saved["extracted_text"]
    assert c.get("/api/documents").json()[0]["data"] == saved["data"]
    assert queue(c, content)["data"] == saved["data"]
    assert calls == ["parse"]
    assert c.get(f'/api/documents/{job["id"]}/download').content == content


def test_parse_network_does_not_block_or_overwrite_persona_edits(
    paused_documents, monkeypatch
):
    c = paused_documents
    started, release = Event(), Event()

    def delayed(self, task, data):
        started.set()
        assert release.wait(5)
        return {"data": {"name": "Parsed Name", "education": "Document education"}}

    monkeypatch.setattr(MockAI, "complete", delayed)
    p = c.post(
        "/api/personas", json={"label": "Existing", "data": {"name": "Manual Name"}}
    ).json()
    job = queue(c)
    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(process_document, job["id"])
        assert started.wait(2)
        try:
            changed = c.put(
                "/api/personas/" + p["id"],
                json={
                    "label": "Edited while parsing",
                    "version": 1,
                    "data": {"name": "My later edit"},
                },
            )
            assert changed.status_code == 200, changed.text
            assert (
                c.get(f'/api/documents/{job["id"]}/status').json()["status"]
                == "processing"
            )
        finally:
            release.set()
        assert future.result(timeout=3)
    assert c.get("/api/personas").json()[0]["data"]["name"] == "My later edit"
    assert (
        c.get(f'/api/documents/{job["id"]}/status').json()["data"]["name"]
        == "Parsed Name"
    )


def test_failure_and_explicit_retry_keep_the_original_file(
    paused_documents, monkeypatch
):
    c = paused_documents
    original = MockAI.complete
    monkeypatch.setattr(
        MockAI,
        "complete",
        lambda *args: (_ for _ in ()).throw(
            HTTPException(502, "AI temporarily unavailable")
        ),
    )
    content = resume_bytes()
    job = queue(c, content)
    process_document(job["id"])
    saved = c.get(f'/api/documents/{job["id"]}/status').json()
    assert saved["status"] == "failed" and saved["data"] is None
    assert c.get(f'/api/documents/{job["id"]}/download').content == content
    monkeypatch.setattr(MockAI, "complete", original)
    retry = c.post(f'/api/documents/{job["id"]}/retry')
    assert retry.status_code == 202 and retry.json()["id"] == job["id"]
    assert c.post(f'/api/documents/{job["id"]}/retry').json()["status"] == "queued"
    process_document(job["id"])
    assert c.get(f'/api/documents/{job["id"]}/status').json()["status"] == "parsed"
    malformed = queue(c, b"not pdf", "bad.pdf")
    process_document(malformed["id"])
    assert (
        c.get(f'/api/documents/{malformed["id"]}/status').json()["status"] == "failed"
    )
    assert (
        c.post("/api/documents/jobs", files={"file": ("bad.exe", b"no")}).status_code
        == 415
    )


def test_parse_recovery_and_workspace_boundaries(paused_documents):
    c = paused_documents
    queued = queue(c, resume_bytes("Queued Person"))
    interrupted = queue(c, resume_bytes("Interrupted Person"))
    with Session() as db:
        db.get(UploadedDocument, interrupted["id"]).status = "processing"
        db.add(Workspace(id="document-other", name="Other"))
        db.flush()
        private = UploadedDocument(
            workspace_id="document-other",
            original_name="private.docx",
            storage_key="missing.docx",
            status="parsed",
            parsed_data={"name": "Private"},
        )
        db.add(private)
        db.commit()
        private_id = private.id
    assert c.get(f"/api/documents/{private_id}/status").status_code == 404
    assert c.post(f"/api/documents/{private_id}/retry").status_code == 404
    assert len(c.get("/api/documents").json()) == 2
    recover_document_jobs()
    state = c.get(f'/api/documents/{interrupted["id"]}/status').json()
    assert state["status"] == "failed" and "may have completed" in state["error"]
    start_document_worker()
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline:
        saved = c.get(f'/api/documents/{queued["id"]}/status').json()
        if saved["status"] == "parsed":
            break
        time.sleep(0.02)
    assert saved["status"] == "parsed"
