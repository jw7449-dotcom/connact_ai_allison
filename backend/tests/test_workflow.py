from io import BytesIO
from docx import Document
from pypdf import PdfWriter
from app.db import Session
from app.models import Workspace, Persona, MatchAssessment, Draft
from app.config import settings


def persona(client):
    r = client.post(
        "/api/personas",
        json={
            "label": "Banking",
            "data": {
                "name": "Taylor Test",
                "sectors": "Investment Banking",
                "target_roles": "Analyst",
                "target_regions": "New York",
            },
        },
    )
    assert r.status_code == 200, r.text
    return r.json()


def search(client, **filters):
    r = client.post("/api/finance/search", json=filters)
    assert r.status_code == 200, r.text
    return r.json()


def test_complete_workflow_and_conflicts(client):
    p = persona(client)
    p = client.put(
        "/api/personas/" + p["id"],
        json={
            "label": "Updated",
            "version": 1,
            "data": {**p["data"], "skills": "Python"},
        },
    ).json()
    assert (
        p["version"] == 2
        and client.get("/api/personas").json()[0]["data"]["skills"] == "Python"
    )
    c = search(client)["items"][0]
    a = client.post(
        "/api/finance/assess", json={"contact_ids": [c["id"]], "persona_id": p["id"]}
    ).json()[0]
    assert (
        "Investment Banking" in a["reason"]
        and a["source_ids"]
        and a["persona_version"] == 2
    )
    assert (
        client.post("/api/contacts/" + c["id"] + "/save").json()["already_saved"]
        is False
    )
    assert (
        client.post("/api/contacts/" + c["id"] + "/save").json()["already_saved"]
        is True
    )
    assert len(client.get("/api/contacts").json()) == 1
    # Search reuses the same provider contact, including concurrent-safe DB unique key.
    assert search(client)["items"][0]["id"] == c["id"]
    d = client.post(
        "/api/drafts",
        json={
            "contact_id": c["id"],
            "persona_id": p["id"],
            "purpose": "learn about your career",
        },
    ).json()
    generated = client.post(
        "/api/drafts/" + d["id"] + "/generate",
        json={"action": "generate", "revision": 1},
    )
    assert generated.status_code == 200, generated.text
    d = generated.json()
    d["subject"] = "Hello {{name}} at {{company}}"
    d["body_html"] = (
        "<p>Dear {{name}},</p><p><strong>Thank you</strong> from {{sender_name}}.</p>"
    )
    saved = client.put("/api/drafts/" + d["id"], json=d)
    assert saved.status_code == 200, saved.text
    assert client.put("/api/drafts/" + d["id"], json=d).status_code == 409
    d = saved.json()
    opened = client.get("/api/drafts/" + d["id"]).json()
    assert opened["body_html"] == d["body_html"]
    prev = client.get("/api/drafts/" + d["id"] + "/preview").json()
    assert c["name"] in prev["subject"] and "Taylor Test" in prev["body_text"]
    assert "{{" not in prev["body_text"] and prev["can_mark_ready"]
    assert not prev["recipient_email"]  # Drafts do not depend on an email address.
    assert (
        client.put("/api/drafts/" + d["id"], json={**d, "status": "ready"}).status_code
        == 200
    )


def test_variables_xss_and_manual_no_persona(client):
    c = client.post(
        "/api/contacts",
        json={"name": "<img src=x onerror=alert(1)>", "company": "Test Firm"},
    ).json()
    d = client.post(
        "/api/drafts",
        json={
            "contact_id": c["id"],
            "subject": "{{company}}",
            "body_html": '<p>{{name}} {{school}} {{sender_name}} {{unknown}}</p><script>alert(1)</script><a href="javascript:alert(1)">x</a>',
        },
    ).json()
    p = client.get("/api/drafts/" + d["id"] + "/preview").json()
    assert p["missing_variables"] == ["school", "sender_name", "unknown"]
    assert (
        "<script>" not in p["body_html"]
        and 'href="javascript:' not in p["body_html"]
        and "<img" not in p["body_html"]
    )
    assert not p["can_mark_ready"]
    assert (
        client.put("/api/drafts/" + d["id"], json={**d, "status": "ready"}).status_code
        == 422
    )
    d["body_html"] = "<p>Hi, I would like to connect.</p>"
    assert client.put("/api/drafts/" + d["id"], json=d).status_code == 200


def test_docx_pdf_and_scan_failure(client):
    doc = Document()
    doc.add_paragraph("Taylor Example")
    doc.add_paragraph("Education")
    doc.add_paragraph("Example University, Finance")
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python, Valuation")
    buf = BytesIO()
    doc.save(buf)
    r = client.post(
        "/api/documents",
        files={
            "file": (
                "resume.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    ).json()
    assert (
        r["status"] == "parsed"
        and r["data"]["education"] == "Example University, Finance"
    )
    p = client.post(
        "/api/personas",
        json={
            "label": "Parsed resume",
            "data": r["data"],
            "document_id": r["document_id"],
        },
    ).json()
    assert p["data"]["skills"] == "Python, Valuation"
    assert (
        client.get("/api/documents/" + r["document_id"] + "/download").content
        == buf.getvalue()
    )
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    pdf = BytesIO()
    writer.write(pdf)
    scan = client.post(
        "/api/documents",
        files={"file": ("scan.pdf", pdf.getvalue(), "application/pdf")},
    ).json()
    assert scan["status"] == "failed" and "Scanned" in scan["error"]
    assert (
        client.post("/api/documents", files={"file": ("bad.exe", b"bad")}).status_code
        == 415
    )
    bad = client.post(
        "/api/documents", files={"file": ("bad.pdf", b"not a pdf")}
    ).json()
    assert bad["status"] == "failed"


def test_search_pagination_filter_limits_and_enrichment(client):
    first = search(client)
    second = search(client, page=2)
    assert first["total"] == 16 and len(second["items"]) == 6
    assert {c["id"] for c in first["items"]}.isdisjoint(
        c["id"] for c in second["items"]
    )
    assert all(
        "Investment Banking" == c["domains"]["finance"]["sector"]
        for c in search(client, sector="Investment Banking")["items"]
    )
    assert search(client, company="nonexistent")["total"] == 0
    assert client.post("/api/finance/search", json={"per_page": 100}).status_code == 422
    c = first["items"][0]
    assert not c["email"]
    r = client.post("/api/contacts/" + c["id"] + "/enrich").json()
    assert r["contact"]["email"].endswith(".example") and not r["cached"]
    assert client.post("/api/contacts/" + c["id"] + "/enrich").json()["cached"]
    r = client.post("/api/contacts/" + c["id"] + "/public-sources").json()
    assert any(s["kind"] == "unverified_lead" for s in r["sources"])
    p = persona(client)
    assert (
        client.post(
            "/api/finance/assess",
            json={"persona_id": p["id"], "contact_ids": [c["id"]] * 6},
        ).status_code
        == 422
    )


def test_workspace_boundaries_and_origin(client):
    with Session() as db:
        db.add(Workspace(id="other", name="Other"))
        db.flush()
        p = Persona(workspace_id="other", label="Private", data={"name": "Private"})
        db.add(p)
        db.commit()
        pid = p.id
    assert client.get("/api/personas").json() == []
    assert client.post("/api/drafts", json={"persona_id": pid}).status_code == 404
    assert (
        client.post(
            "/api/personas",
            headers={"Origin": "https://evil.example"},
            json={"label": "x", "data": {}},
        ).status_code
        == 403
    )


def test_live_missing_credentials_never_fallback(client, monkeypatch):
    monkeypatch.setattr(settings, "people_mode", "live")
    monkeypatch.setattr(settings, "serpapi_api_key", "")
    r = client.post("/api/finance/search", json={})
    assert r.status_code == 503 and "Live" in r.text


def test_persona_revision_assessment_cache(client):
    p = persona(client)
    c = search(client)["items"][0]
    req = {"persona_id": p["id"], "contact_ids": [c["id"]]}
    a = client.post("/api/finance/assess", json=req).json()[0]
    assert client.post("/api/finance/assess", json=req).json()[0]["id"] == a["id"]
    client.put(
        "/api/personas/" + p["id"],
        json={
            "label": p["label"],
            "data": {**p["data"], "target_roles": "Portfolio Manager"},
            "version": 1,
        },
    )
    b = client.post("/api/finance/assess", json=req).json()[0]
    assert b["id"] != a["id"] and b["persona_revision_id"] != a["persona_revision_id"]


def test_shortening_preserves_custom_draft(client):
    d = client.post(
        "/api/drafts",
        json={
            "subject": "My custom subject",
            "body_html": "<p>Hi there,</p><p>Optional introduction.</p><p>My unique request.</p><p>Thanks, Taylor</p>",
        },
    ).json()
    r = client.post(
        "/api/drafts/" + d["id"] + "/generate",
        json={"action": "shorten", "revision": 1},
    ).json()
    assert (
        r["subject"] == "My custom subject"
        and "My unique request" in r["body_html"]
        and len(r["body_html"]) < len(d["body_html"])
    )


def test_text_pdf_extraction(client):
    from pathlib import Path

    content = (
        Path(__file__).resolve().parents[2] / "demo" / "sample-resume.pdf"
    ).read_bytes()
    r = client.post(
        "/api/documents", files={"file": ("resume.pdf", content, "application/pdf")}
    ).json()
    assert (
        r["status"] == "parsed"
        and r["data"]["name"] == "Alex Morgan"
        and "New York University" in r["data"]["education"]
    )


def test_background_changes_invalidate_reviewed_drafts(client):
    p = persona(client)
    c = client.post(
        "/api/contacts", json={"name": "Jane", "school": "Test School"}
    ).json()
    d = client.post(
        "/api/drafts",
        json={
            "contact_id": c["id"],
            "persona_id": p["id"],
            "subject": "Hello",
            "body_html": "<p>{{school}}</p>",
            "status": "ready",
        },
    ).json()
    assert d["status"] == "ready"
    client.put("/api/contacts/" + c["id"], json={"name": "Jane", "school": ""})
    changed = client.get("/api/drafts/" + d["id"]).json()
    assert (
        changed["status"] == "draft"
        and changed["revision"] == 2
        and not changed["preview"]["can_mark_ready"]
    )
    assert client.put("/api/drafts/" + d["id"], json=d).status_code == 409
    client.put(
        "/api/personas/" + p["id"],
        json={"label": p["label"], "data": p["data"], "version": 1},
    )
    changed = client.get("/api/drafts/" + d["id"]).json()
    assert changed["preview"]["persona_changed"] and changed["revision"] == 3


def test_postgres_concurrent_draft_updates(client):
    import pytest
    from concurrent.futures import ThreadPoolExecutor
    from app.db import engine

    if engine.dialect.name != "postgresql":
        pytest.skip("Requires PostgreSQL row locks")
    d = client.post("/api/drafts", json={"subject": "Original"}).json()

    def write(subject):
        return client.put(
            "/api/drafts/" + d["id"], json={**d, "subject": subject}
        ).status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        statuses = list(pool.map(write, ["Writer A", "Writer B"]))
    assert sorted(statuses) == [200, 409]
    assert client.get("/api/drafts/" + d["id"]).json()["revision"] == 2
